"""Extract frozen BNS tables and construct an auditable regional panel.

Run from any directory: python build_dataset.py
No network calls and no fabricated/interpolated observations.
"""
from pathlib import Path
import re, json, hashlib
import numpy as np
import pandas as pd
import openpyxl

ROOT=Path(__file__).resolve().parent
RAW=ROOT/'data'/'raw'; PROCESSED=ROOT/'data'/'processed'
REGIONS={100000000:('Abay','Абай'),110000000:('Akmola','Акмолинская'),150000000:('Aktobe','Актюбинская'),190000000:('Almaty region','Алматинская'),230000000:('Atyrau','Атырауская'),270000000:('West Kazakhstan','Западно-Казахстанская'),310000000:('Zhambyl','Жамбылская'),330000000:('Zhetisu','Жетісу'),350000000:('Karaganda','Карагандинская'),390000000:('Kostanay','Костанайская'),430000000:('Kyzylorda','Кызылординская'),470000000:('Mangystau','Мангистауская'),550000000:('Pavlodar','Павлодарская'),590000000:('North Kazakhstan','Северо-Казахстанская'),610000000:('Turkistan','Туркестанская'),620000000:('Ulytau','Ұлытау'),630000000:('East Kazakhstan','Восточно-Казахстанская'),710000000:('Astana','Астана'),750000000:('Almaty city','город Алматы'),790000000:('Shymkent','Шымкент')}
FEATURES=['income_ratio_lag1','wage_ratio_lag1','unemployment_lag1','urban_share_lag1','log_population_lag1']
BOUNDARY_2022={100000000,190000000,330000000,350000000,620000000,630000000}
BOUNDARY_2018={610000000,790000000}

def normalise(value):
    s=re.sub(r'\d+\)','',str(value or '')).lower().strip()
    s=s.replace('область','').replace('город','').replace('г.','').replace(' ','')
    return s

def number(value):
    if isinstance(value,(int,float)): return float(value)
    # Dashes in historical cells do not create zero-valued socioeconomic data.
    return np.nan

def year_columns(ws,row):
    cols={}
    for cells in ws.iter_rows(min_row=row,max_row=row):
        for c in cells:
            m=re.match(r'^(20\d{2})',str(c.value))
            if m: cols[int(m.group(1))]=c.column
    return cols

def build_dataset():
    PROCESSED.mkdir(parents=True,exist_ok=True)
    manifest=json.loads((ROOT/'source_manifest.json').read_text(encoding='utf-8'))
    for src in manifest:
        assert hashlib.sha256((ROOT/src['file']).read_bytes()).hexdigest()==src['sha256'],src['id']
    books={p.stem:openpyxl.load_workbook(p,data_only=True) for p in RAW.glob('*.xlsx')}
    lineage=[]; extracted={}
    metadata={name:{str(ws.cell(r,1).value):str(ws.cell(r,2).value) for r in range(1,ws.max_row+1) if ws.cell(r,1).value} for name,wb in books.items() for ws in [wb['Метаданные']]}
    (PROCESSED/'source_metadata.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')

    def extract(name,sheet,header_row,label_col,start,end,variable,code_col=None,skip_rows=()):
        ws=books[name][sheet]; cols=year_columns(ws,header_row)
        rows={}
        for r in range(start,end+1):
            if r in skip_rows: continue
            if code_col:
                rawcode=ws.cell(r,code_col).value
                try: key=int(rawcode)
                except (TypeError,ValueError): key=0 if normalise(ws.cell(r,label_col).value)==normalise('Республика Казахстан') else None
            else:
                label=normalise(ws.cell(r,label_col).value)
                if label==normalise('Республика Казахстан'): key=0
                else:
                    found=[k for k,(eng,ru) in REGIONS.items() if normalise(ru)==label]
                    key=found[0] if found else None
            if key is None or (key!=0 and key not in REGIONS): continue
            assert key not in rows,(name,variable,key,r)
            rows[key]=r
            for year,col in cols.items():
                if not 2014<=year<=2025: continue
                cell=ws.cell(r,col); value=number(cell.value)
                extracted[(key,year,variable)]=value
                lineage.append(dict(source_id=name,sheet=sheet,cell=cell.coordinate,kato=key,year=year,variable=variable,raw_value=cell.value,value=value))
        return rows

    extract('migration_net','Показатель',5,2,7,27,'migration_net',1)
    extract('migration_in','Показатель',6,2,8,28,'migration_in',1)
    extract('migration_out','Показатель',6,2,8,28,'migration_out',1)
    extract('population','2009-2025',7,2,9,29,'population',1)
    extract('population','2009-2025',7,2,75,95,'urban_population',1)
    extract('income','Показатель',4,2,5,26,'income',1)
    # Region totals, both sexes, not similarly named subregional city records.
    extract('unemployment','Показатель',4,1,6,294,'unemployment',skip_rows=(182,))
    # The first Shymkent record is the old subregional city; its label normalises
    # identically to the new republic-level city but is not a study unit.
    # Wages contain monthly/quarterly columns and a separate annual value.
    ws=books['wages']['квартал']; wage_year_cols={}; current_year=None
    for col in range(2,ws.max_column+1):
        head=ws.cell(3,col).value
        if isinstance(head,int) and 2000<=head<=2030: current_year=head
        if str(ws.cell(4,col).value).strip()=='за год' and 2014<=current_year<=2025:
            wage_year_cols[current_year]=col
    for row in range(6,28):
        label=normalise(ws.cell(row,1).value)
        matches=[k for k,(_,ru) in REGIONS.items() if normalise(ru)==label]
        key=0 if label==normalise('Республика Казахстан') else (matches[0] if matches else None)
        if key is None: continue
        for year,col in wage_year_cols.items():
            cell=ws.cell(row,col); value=number(cell.value)
            extracted[(key,year,'wage')]=value
            lineage.append(dict(source_id='wages',sheet=ws.title,cell=cell.coordinate,kato=key,year=year,variable='wage',raw_value=cell.value,value=value))
    pd.DataFrame(lineage).to_csv(PROCESSED/'cell_lineage.csv',index=False,encoding='utf-8-sig')
    variables=['migration_in','migration_out','migration_net','population','urban_population','income','wage','unemployment']
    records=[]
    for year in range(2014,2026):
        for kato,(eng,ru) in REGIONS.items():
            row=dict(kato=kato,region=eng,region_ru=ru,year=year)
            for var in variables: row[var]=extracted.get((kato,year,var),np.nan)
            row['income_national']=extracted.get((0,year,'income'),np.nan)
            row['wage_national']=extracted.get((0,year,'wage'),np.nan)
            row['income_ratio']=row['income']/row['income_national']
            row['wage_ratio']=row['wage']/row['wage_national']
            row['urban_share']=row['urban_population']/row['population']
            row['log_population']=np.log(row['population'])
            row['economic_geometry_comparable']=(year>=2022 if kato in BOUNDARY_2022 else (year>=2018 if kato in BOUNDARY_2018 else True))
            records.append(row)
    panel=pd.DataFrame(records).sort_values(['year','kato']).reset_index(drop=True)
    assert len(panel)==240
    for year,group in panel.groupby('year'):
        assert group.migration_net.sum()==0,(year,'net balance does not sum to zero')
        assert (group.migration_in-group.migration_out==group.migration_net).all(),year
        assert ((group.urban_share>=0)&(group.urban_share<=1)).all(),year
    panel.to_csv(PROCESSED/'regional_panel.csv',index=False,encoding='utf-8-sig')

    exclusions=[]; modelling=[]; indexed=panel.set_index(['kato','year'])
    feature_bases=['income_ratio','wage_ratio','unemployment','urban_share','log_population']
    for rec in panel.loc[panel.year>=2015].to_dict('records'):
        kato=rec['kato']; year=rec['year']; fy=year-1
        reason=None
        if kato in BOUNDARY_2022 and fy<2022:
            reason='2022 territorial split: socioeconomic geometry not certified comparable before 2022'
        elif kato in BOUNDARY_2018 and fy<2018:
            reason='2018 territorial split: no comparable pre-2018 socioeconomic series'
        elif pd.isna(rec['migration_net']): reason='missing target'
        elif rec['migration_net']==0: reason='zero net balance excluded from positive/negative task'
        prev=indexed.loc[(kato,fy)]
        features={var+'_lag1':prev[var] for var in feature_bases}
        if reason is None and any(pd.isna(v) for v in features.values()): reason='missing required lagged socioeconomic feature'
        if reason:
            exclusions.append(dict(kato=kato,region=rec['region'],year=year,reason=reason)); continue
        rec.update(features)
        rec.update(feature_year=fy,target=int(rec['migration_net']>0),previous_net=prev.migration_net,
                   persistence_prediction=int(prev.migration_net>0),split='test' if year>=2024 else 'development')
        modelling.append(rec)
    dataset=pd.DataFrame(modelling).sort_values(['year','kato']).reset_index(drop=True)
    dataset.to_csv(PROCESSED/'model_dataset.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame(exclusions).to_csv(PROCESSED/'excluded_rows.csv',index=False,encoding='utf-8-sig')
    summary=dataset.groupby('year').agg(rows=('target','size'),positive=('target','sum'))
    summary['negative']=summary.rows-summary.positive
    summary.to_csv(PROCESSED/'class_counts_by_year.csv',encoding='utf-8-sig')
    missing=panel[variables].isna().sum().rename('raw_panel_missing').to_frame()
    missing.to_csv(PROCESSED/'missingness.csv',encoding='utf-8-sig')
    dictionary=[
        ('kato','Official current regional code','identifier','Not a model feature'),
        ('year','Migration outcome year','calendar year','2015-2025 in modelling data'),
        ('feature_year','Previous calendar year','calendar year','year minus 1'),
        ('migration_in','Internal arrivals, all population','persons','BNS 613301; verification only'),
        ('migration_out','Internal departures, all population','persons','BNS 613302; verification only'),
        ('migration_net','Internal arrivals minus internal departures','persons','BNS 613303; target construction only'),
        ('target','Positive migration balance = 1; negative = 0','binary','Exact zeros excluded; not observed in final dataset'),
        ('population','Average annual total population','persons','BNS 611104'),
        ('urban_population','Average annual urban population','persons','BNS 611104 urban block'),
        ('income','Nominal per capita monetary income, monthly level in annual series','KZT/person/month','BNS 641101; annual series, not sum of months'),
        ('wage','Average monthly nominal wage, official annual value','KZT/employee/month','Annual column za god; not sum of quarterly wages'),
        ('unemployment','Unemployment rate, both sexes, region total','percent','BNS unemployment annual series'),
        ('income_ratio_lag1','Previous-year regional income / same-year national income','ratio','1 means national average; not a regional cost-of-living adjustment'),
        ('wage_ratio_lag1','Previous-year regional wage / same-year national wage','ratio','1 means national average'),
        ('unemployment_lag1','Previous-year unemployment rate','percent','Both sexes total'),
        ('urban_share_lag1','Previous-year urban population / total population','fraction 0-1','Same region, same year, same source'),
        ('log_population_lag1','Natural logarithm of previous-year average population','log persons','Controls scale without using region identity'),
        ('previous_net','Previous-year internal net migration','persons','Only persistence benchmark; excluded from socioeconomic ML inputs'),
        ('split','Development or final held-out test','category','Development <=2023; test 2024-2025')
    ]
    pd.DataFrame(dictionary,columns=['variable','definition','unit','source_or_rule']).to_csv(PROCESSED/'data_dictionary.csv',index=False,encoding='utf-8-sig')
    summary_meta=dict(raw_rows=len(panel),candidate_rows=220,model_rows=len(dataset),excluded_rows=len(exclusions),
                      development_rows=int((dataset.year<=2023).sum()),test_rows=int((dataset.year>=2024).sum()),
                      positive=int(dataset.target.sum()),negative=int((dataset.target==0).sum()),
                      features=FEATURES,feature_missing=int(dataset[FEATURES].isna().sum().sum()),
                      territory_strategy='Conservative unbalanced panel; all 20 current units from outcome year 2023; split-affected regions omitted earlier, not imputed')
    (PROCESSED/'dataset_summary.json').write_text(json.dumps(summary_meta,indent=2),encoding='utf-8')
    print(json.dumps(summary_meta,indent=2)); print(summary.to_string())
    for wb in books.values(): wb.close()
    return dataset

def main():
    build_dataset()

if __name__=='__main__': main()
